"""Run the standard pipeline and create the Excel analyst workbook."""
from __future__ import annotations

import csv
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "artifacts" / "rating_drift_dashboard.xlsx"


def _read_metrics() -> list[dict[str, str]]:
    csv_path = ROOT / "data" / "processed" / "latest_credit_metrics.csv"
    with csv_path.open("r", newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _read_companies() -> dict[str, dict[str, str]]:
    csv_path = ROOT / "data" / "raw" / "companies.csv"
    with csv_path.open("r", newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    return {row["company_id"]: row for row in rows}


def export_workbook(output_path: Path | str = OUTPUT_PATH) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    rows = _read_metrics()
    companies = _read_companies()
    ordered = sorted(rows, key=lambda r: float(r["risk_score"]), reverse=True)

    wb = Workbook()
    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"
    ws_dashboard.sheet_view.showGridLines = False

    ws_dashboard["A1"] = "Rating Drift Detector"
    ws_dashboard["A1"].font = Font(bold=True, size=16, color="1F4E78")
    ws_dashboard["A2"] = "Credit-surveillance summary | Sample data | Latest quarter: 2025-Q4"
    ws_dashboard["A2"].font = Font(italic=True, color="666666")

    summary = [
        ["Metric", "Value"],
        ["Companies monitored", len(rows)],
        ["High-risk companies", sum(1 for r in rows if r["risk_band"] == "High")],
        ["Average risk score", round(sum(float(r["risk_score"]) for r in rows) / len(rows), 1) if rows else 0.0],
    ]
    for row in summary:
        ws_dashboard.append(row)

    for cell in ws_dashboard["A5:B8"][0]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")

    headings_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws_dashboard.iter_rows(min_row=5, max_row=8, min_col=1, max_col=2):
        for cell in row:
            if cell.row == 5:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(bold=True, color="FFFFFF")
            else:
                cell.fill = headings_fill
                cell.font = Font(bold=True)

    ws_dashboard["D5"] = "Company"
    ws_dashboard["E5"] = "Risk score"
    ws_dashboard["F5"] = "Risk band"
    ws_dashboard["G5"] = "Key credit signal"
    ws_dashboard["H5"] = "Analyst action"
    for cell in ws_dashboard[5]:
        if cell.column in ("D", "E", "F", "G", "H"):
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(bold=True, color="FFFFFF")

    for idx, row in enumerate(ordered, start=6):
        company = companies.get(row["company_id"], {})
        ws_dashboard[f"D{idx}"] = company.get("company_name", row["company_id"])
        ws_dashboard[f"E{idx}"] = int(float(row["risk_score"]))
        ws_dashboard[f"F{idx}"] = row["risk_band"]
        ws_dashboard[f"G{idx}"] = row["flags"]
        ws_dashboard[f"H{idx}"] = "Review at next credit committee" if row["risk_band"] == "High" else "Add to quarterly watchlist" if row["risk_band"] == "Moderate" else "Continue routine monitoring"

    for col in ["A", "B", "D", "E", "F", "G", "H"]:
        ws_dashboard.column_dimensions[col].width = 20
    ws_dashboard.column_dimensions["C"].width = 8
    ws_dashboard.column_dimensions["G"].width = 35
    ws_dashboard.column_dimensions["H"].width = 30

    ws_watchlist = wb.create_sheet("Watchlist")
    ws_watchlist.append([
        "Company",
        "Sector",
        "Quarter",
        "Revenue growth",
        "Debt / EBITDA",
        "Interest coverage",
        "Current ratio",
        "Receivable days",
        "Cash conversion",
        "Risk score",
        "Risk band",
        "Credit flags",
    ])
    for row in ordered:
        company = companies.get(row["company_id"], {})
        ws_watchlist.append([
            company.get("company_name", row["company_id"]),
            company.get("sector", ""),
            row["quarter"],
            float(row["revenue_growth"]),
            float(row["debt_to_ebitda"]),
            float(row["interest_coverage"]),
            float(row["current_ratio"]),
            float(row["receivable_days"]),
            float(row["cash_conversion"]),
            int(float(row["risk_score"])),
            row["risk_band"],
            row["flags"],
        ])
    for cell in ws_watchlist[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")

    ws_metrics = wb.create_sheet("Metrics data")
    ws_metrics.append([
        "Company ID",
        "Quarter",
        "Revenue growth",
        "Debt / EBITDA",
        "Interest coverage",
        "Current ratio",
        "Receivable days",
        "Cash conversion",
        "Risk score",
        "Risk band",
        "Credit flags",
    ])
    for row in rows:
        ws_metrics.append([
            row["company_id"],
            row["quarter"],
            float(row["revenue_growth"]),
            float(row["debt_to_ebitda"]),
            float(row["interest_coverage"]),
            float(row["current_ratio"]),
            float(row["receivable_days"]),
            float(row["cash_conversion"]),
            int(float(row["risk_score"])),
            row["risk_band"],
            row["flags"],
        ])
    for cell in ws_metrics[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")

    ws_readme = wb.create_sheet("Read me")
    ws_readme["A1"] = "Workbook guide"
    ws_readme["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws_readme["A3"] = "Section"
    ws_readme["B3"] = "Purpose"
    ws_readme["A3"].fill = PatternFill("solid", fgColor="1F4E78")
    ws_readme["B3"].fill = PatternFill("solid", fgColor="1F4E78")
    ws_readme["A3"].font = Font(bold=True, color="FFFFFF")
    ws_readme["B3"].font = Font(bold=True, color="FFFFFF")

    guide = [
        ("Dashboard", "Credit-surveillance summary and prioritised analyst actions."),
        ("Watchlist", "Peer metrics, risk flags and company-level review view."),
        ("Metrics data", "Latest output from the Python scoring pipeline."),
        ("Method", "The score adds points for leverage, weak interest coverage, tight liquidity, stretched receivables and weak cash conversion."),
        ("Limitation", "A triage tool for learning. It is not a rating opinion or a substitute for analyst judgement."),
    ]
    for idx, (section, purpose) in enumerate(guide, start=4):
        ws_readme[f"A{idx}"] = section
        ws_readme[f"B{idx}"] = purpose
        ws_readme[f"B{idx}"].alignment = Alignment(wrap_text=True)

    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
        ws_watchlist.column_dimensions[col].width = 16
    ws_watchlist.column_dimensions["L"].width = 35
    ws_metrics.column_dimensions["A"].width = 14
    ws_metrics.column_dimensions["L"].width = 35
    ws_readme.column_dimensions["A"].width = 20
    ws_readme.column_dimensions["B"].width = 90

    wb.save(output)
    print(f"Workbook created: {output}")
    return output


if __name__ == "__main__":
    subprocess.run([sys.executable, "-m", "src.pipeline"], cwd=ROOT, check=True)
    export_workbook()
